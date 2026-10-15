"""Read a registered workbook by input ID, with bounded rows and saved formula caches."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import stat
import subprocess

MAX_BYTES = 20 << 20
MAX_OUTPUT = 24000


class SourceError(ValueError):
    pass


def require(condition, code):
    if not condition:
        raise SourceError(code)


def registered_bytes(input_id):
    reply = subprocess.run(['/usr/local/bin/cc-connect-file', 'work-context'], input=b'{}',
                           capture_output=True, timeout=30, check=True)
    context = json.loads(reply.stdout)
    require(context.get('enabled') is True, 'file_work_disabled')
    selected = [item for item in context['inputs'] if item['id'] == input_id]
    require(len(selected) == 1, 'input_id_not_unique')
    item = selected[0]
    path = Path(item['path'])
    output = Path(context['output_dir'])
    require(path.is_absolute() and output.is_absolute() and output.name == 'outputs'
            and path.parent == output.parent / 'inputs', 'input_outside_work')
    require(path.suffix.lower() == '.xlsx', 'expected_xlsx')
    parents = {}
    for parent in path.parents:
        info = parent.lstat()
        require(stat.S_ISDIR(info.st_mode) and not parent.is_symlink(), 'unsafe_input_parent')
        parents[parent] = (info.st_dev, info.st_ino)
    before = path.lstat()
    require(stat.S_ISREG(before.st_mode) and not path.is_symlink() and before.st_nlink == 1,
            'unsafe_input_file')
    require(0 < before.st_size <= MAX_BYTES, 'input_size_limit')
    fd = os.open(path, os.O_RDONLY | getattr(os, 'O_BINARY', 0)
                 | getattr(os, 'O_NOFOLLOW', 0) | getattr(os, 'O_NONBLOCK', 0))
    with os.fdopen(fd, 'rb') as stream:
        opened = os.fstat(stream.fileno())
        require((before.st_dev, before.st_ino, before.st_mtime_ns, before.st_size)
                == (opened.st_dev, opened.st_ino, opened.st_mtime_ns, opened.st_size), 'input_changed')
        data = stream.read(MAX_BYTES + 1)
        after = os.fstat(stream.fileno())
    current = path.lstat()
    identity = lambda s: (s.st_dev, s.st_ino, s.st_size, s.st_mtime_ns, s.st_nlink)
    require(identity(before) == identity(after) == identity(current), 'input_changed')
    # Windows lstat/fstat expose different ctime meanings; compare within each API.
    require(before.st_ctime_ns == current.st_ctime_ns and opened.st_ctime_ns == after.st_ctime_ns,
            'input_changed')
    require(len(data) == before.st_size, 'input_changed')
    for parent, expected in parents.items():
        info = parent.lstat()
        require(stat.S_ISDIR(info.st_mode) and (info.st_dev, info.st_ino) == expected,
                'input_parent_changed')
    require(hashlib.sha256(data).hexdigest() == item['sha256'], 'input_hash_changed')
    checked = subprocess.run(['/usr/local/bin/cc-connect-anc', 'file-check', 'source.xlsx'],
                             input=data, capture_output=True, timeout=30, check=True)
    require(checked.stdout == b'OFFICE_OK_V1\n', 'office_validation_failed')
    return item, data


def read_rows(input_id, sheet_name, start_row, max_rows):
    from openpyxl import load_workbook
    item, data = registered_bytes(input_id)
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    cached = None
    try:
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        require(sheet_name is None or sheet_name in book.sheetnames, 'sheet_not_found')
        sheet = book[sheet_name or book.sheetnames[0]]
        require(sheet.max_column <= 512, 'worksheet_width_limit')
        require(start_row <= max(sheet.max_row, 1), 'row_out_of_range')
        end_row = min(sheet.max_row, start_row + max_rows - 1)
        rows = []
        saved_rows = cached[sheet.title].iter_rows(min_row=start_row, max_row=end_row)
        for number, (row, saved) in enumerate(zip(sheet.iter_rows(min_row=start_row, max_row=end_row), saved_rows), start_row):
            cells = {cell.coordinate: {'formula': cell.value, 'cached': value.value}
                     if cell.data_type == 'f' else cell.value
                     for cell, value in zip(row, saved) if cell.value is not None}
            rows.append({'row': number, 'cells': cells})
        return {'status': 'read', 'source': {key: item[key] for key in ('id', 'name', 'path', 'sha256')},
                'worksheets': book.sheetnames, 'sheet': sheet.title, 'start_row': start_row,
                'end_row': end_row, 'last_row': sheet.max_row, 'rows': rows,
                'next_row': end_row + 1 if end_row < sheet.max_row else None,
                'scope': 'Selected sheet rows only; repeat for remaining rows and sheets. Formula caches are saved values, not a recalculation.'}
    finally:
        book.close()
        if cached is not None:
            cached.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-id', required=True)
    parser.add_argument('--sheet')
    parser.add_argument('--start-row', type=int, default=1)
    parser.add_argument('--max-rows', type=int, default=50)
    args = parser.parse_args(argv)
    try:
        require(args.start_row > 0 and 1 <= args.max_rows <= 200, 'invalid_row_range')
        value = read_rows(args.input_id, args.sheet, args.start_row, args.max_rows)
        output = json.dumps(value, ensure_ascii=False, default=str)
        require(len(output) <= MAX_OUTPUT, 'readback_too_large_reduce_rows')
        print(output)
        return 0
    except Exception as exc:
        # Host/validator stderr may contain private context; do not echo it.
        code = str(exc) if isinstance(exc, SourceError) else type(exc).__name__
        print(json.dumps({'status': 'unread', 'error': code}))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
