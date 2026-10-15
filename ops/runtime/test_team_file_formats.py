"""Fictional packages exercise parser decisions and editable media round trips."""
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import unittest

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject

from team_file_pdf import validate_pdf


def pdf_bytes(modify=lambda writer: None):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    modify(writer)
    data = io.BytesIO()
    writer.write(data)
    return data.getvalue()


class FileFormatsTests(unittest.TestCase):
    def test_editable_slide_and_poster_material(self):
        for image_format in ('PNG', 'JPEG'):
            picture = io.BytesIO()
            Image.new('RGB', (80, 60), '#ab7858').save(picture, format=image_format)
            picture.seek(0)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = '虚构婚宴方案 · v1'
            slide.shapes.add_picture(picture, Inches(1), Inches(2))
            data = io.BytesIO()
            deck.save(data)
            data.seek(0)
            saved = Presentation(data)
            self.assertEqual(saved.slides[0].shapes[0].text, '虚构婚宴方案 · v1')
            self.assertEqual(len(saved.slides[0].shapes), 2)
            saved.slides[0].shapes[0].text = '虚构婚宴方案 · v2'
            self.assertNotEqual(saved.slides[0].shapes[0].text, deck.slides[0].shapes[0].text)

    def test_pdf_structure_and_active_content(self):
        self.assertEqual(validate_pdf(pdf_bytes()), 1)
        self.assertEqual(validate_pdf(pdf_bytes(lambda w: w._root_object.update({
            NameObject('/OpenAction'): ArrayObject([w.pages[0].indirect_reference, NameObject('/Fit')])}))), 1)
        attacks = [
            lambda w: w.add_js('app.alert("fictional")'),
            lambda w: w.add_attachment('fictional.txt', b'fictional'),
            lambda w: w.encrypt('fictional-password'),
            lambda w: w._root_object.update({NameObject('/OpenAction'): DictionaryObject({
                NameObject('/S'): NameObject('/URI'), NameObject('/URI'): TextStringObject('https://example.invalid')})}),
            lambda w: w._root_object.update({NameObject('/AcroForm'): DictionaryObject()}),
        ]
        for modify in attacks:
            with self.subTest(modify=attacks.index(modify)), self.assertRaises(ValueError):
                validate_pdf(pdf_bytes(modify))
        for data in (b'not PDF', pdf_bytes()[:100], b'%PDF-1.7\ninvalid'):
            with self.assertRaises(Exception):
                validate_pdf(data)

    @unittest.skipUnless(os.name == 'posix', 'production parser resource limits require POSIX')
    def test_bounded_stdin_contract(self):
        command = [sys.executable, '-I', str(Path(__file__).with_name('team_file_pdf.py'))]
        good = subprocess.run(command, input=pdf_bytes(), capture_output=True, timeout=15)
        self.assertEqual((good.returncode, good.stdout), (0, b'PDF_OK_V1\n'))
        bad = subprocess.run(command, input=b'private fictional malformed bytes', capture_output=True, timeout=15)
        self.assertNotEqual(bad.returncode, 0)
        self.assertEqual(bad.stdout, b'')
        self.assertNotIn(b'private fictional', bad.stderr)

    @unittest.skipUnless(os.name == 'posix', 'production parser resource limits require POSIX')
    def test_pdf_read_returns_every_saved_page_not_a_check_claim(self):
        command = [sys.executable, '-I', str(Path(__file__).with_name('team_file_pdf.py')), '--read-content']
        data = pdf_bytes(lambda w: w.add_blank_page(width=612, height=792))
        good = subprocess.run(command, input=data, capture_output=True, timeout=15)
        self.assertEqual(good.returncode, 0)
        self.assertEqual(json.loads(good.stdout), {'format': 'pdf_text_layer', 'parts': [
            {'name': 'page-1', 'text': ''}, {'name': 'page-2', 'text': ''}]})
        bad = subprocess.run(command, input=b'private malformed content', capture_output=True, timeout=15)
        self.assertNotEqual(bad.returncode, 0)
        self.assertEqual(bad.stdout, b'')


if __name__ == '__main__':
    unittest.main()
