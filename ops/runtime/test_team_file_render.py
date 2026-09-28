"""Offline rendering only: fictional sources, no models or external messages."""
import hashlib
import io
import json
import os
import re
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
import xml.etree.ElementTree as ET
from contextlib import redirect_stderr
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from openpyxl import Workbook, load_workbook

from team_file_render import check_xlsx_result, main, office_check, prepare_paths, render, run, xlsx_cells


def deny_unix_sockets():
    """Match the deployed Linux sandbox's AF_UNIX denial in the actual worker."""
    import ctypes as c
    import socket
    assert os.geteuid() != 0
    class Arg(c.Structure):
        _fields_ = [('arg', c.c_uint), ('op', c.c_uint), ('a', c.c_uint64), ('b', c.c_uint64)]
    lib = c.CDLL('libseccomp.so.2')
    lib.seccomp_init.restype = c.c_void_p
    lib.seccomp_rule_add_array.argtypes = [c.c_void_p, c.c_uint32, c.c_int, c.c_uint, c.POINTER(Arg)]
    lib.seccomp_load.argtypes = [c.c_void_p]
    lib.seccomp_release.argtypes = [c.c_void_p]
    lib.seccomp_syscall_resolve_name.argtypes = [c.c_char_p]
    context = lib.seccomp_init(0x7fff0000)  # SCMP_ACT_ALLOW
    assert context
    try:
        # SCMP_ACT_ERRNO(EPERM), argument 0 == AF_UNIX (SCMP_CMP_EQ).
        syscall = lib.seccomp_syscall_resolve_name(b'socket')
        assert syscall >= 0
        assert lib.seccomp_rule_add_array(context, 0x50001, syscall, 1, c.byref(Arg(0, 4, socket.AF_UNIX, 0))) == 0
        assert lib.seccomp_load(context) == 0
    finally:
        lib.seccomp_release(context)
    try:
        with socket.socket(socket.AF_UNIX):
            raise AssertionError('AF_UNIX unexpectedly allowed')
    except PermissionError:
        pass


def add_field(paragraph, code):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), code)
    run_node = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '0'
    run_node.append(text)
    field.append(run_node)
    paragraph._p.append(field)


def fictional_template(path, extra=False):
    doc = Document()
    for style in ('Normal', 'Heading 1', 'Heading 2'):
        doc.styles[style].font.name = 'Noto Sans CJK SC'
        doc.styles[style]._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    doc.styles['Normal'].font.size = Pt(11)
    doc.sections[0].header.paragraphs[0].text = '虚构婚宴演练资料 v1'
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run('Page ')
    add_field(footer, 'PAGE')
    footer.add_run(' of ')
    add_field(footer, 'NUMPAGES')
    doc.add_paragraph('目录')
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_page_break()
    doc.add_heading('第一章 客户需求', level=1)
    doc.add_paragraph('虚构客户，22 桌，220 人。来源：模拟套餐资料 v1。')
    if extra:
        for i in range(3):
            doc.add_page_break()
            doc.add_paragraph(f'新增演练说明 {i+1}，非真实业务。')
    doc.add_page_break()
    doc.add_heading('第二章 执行安排', level=1)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ('序号', '工作', '状态')):
        cell.text = text
    for i in range(65):
        for cell, text in zip(table.add_row().cells, (str(i+1), '虚构工作条目', '待确认')):
            cell.text = text
    doc.save(path)


def fictional_budget(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '虚构婚宴预算'
    for row in (
        ('虚构婚宴预算 v2', '数量', '单价', '金额'),
        ('A 套餐', 22, 4080, '=B2*C2'),
        ('酒水', 132, 30, '=B3*C3'),
        ('固定服务及主持', None, None, 5000),
        ('基础布置', None, None, 5000),
        ('合计', None, None, '=SUM(D2:D5)'),
        ('预算余量', None, None, '=110000-D6'),
        ('预算判断', None, None, '=IF(D6<=110000,"预算内","超预算")'),
        ('误存为文本的公式', '=SUM(D2:D5)'),
        ('正常说明', '= 表示预算口径说明，不是公式'),
        ('跨表来源价格', None, None, "='依据 v2'!B1"),
        ('空字符串结果', None, None, '=IF(D6>0,"","")'),
    ):
        sheet.append(row)
    sheet['B9'].data_type = 's'
    sheet['B10'].data_type = 's'
    for name, width in (('A', 24), ('B', 43), ('C', 14), ('D', 18)):
        sheet.column_dimensions[name].width = width
    for row in sheet.iter_rows(min_row=2, max_row=7, min_col=4, max_col=4):
        row[0].number_format = '#,##0'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = 'A1:D12'
    reference = workbook.create_sheet('依据 v2')
    reference.append(('虚构 Wiki v2 A 套餐单价', 4080))
    reference.column_dimensions['A'].width = 32
    workbook.save(path)


def saved_budget_formulas(formulas, *, cached='6', extra_sheets=()):
    """Two local sheets with explicit saved caches, without evaluating formulas."""
    book = Workbook()
    book.active.title = '预算总览'
    book.create_sheet('预算明细')
    for title in extra_sheets:
        book.create_sheet(title)['A1'] = 1
    book['预算总览']['B5'] = 2
    book['预算明细']['C2'] = 3
    for address, formula in formulas.items():
        title, coordinate = address.rsplit('!', 1)
        book[title][coordinate] = formula
    raw = io.BytesIO()
    book.save(raw)
    result = io.BytesIO()
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(raw) as original, zipfile.ZipFile(result, 'w') as target:
        for name in original.namelist():
            value = original.read(name)
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                xml = ET.fromstring(value)
                for cell in xml.iter(ns + 'c'):
                    if cell.find(ns + 'f') is not None:
                        cell.find(ns + 'v').text = cached
                value = ET.tostring(xml)
            target.writestr(name, value)
    return result.getvalue()


class RenderBoundaryTests(unittest.TestCase):
    def test_optional_local_sheet_quotes_survive_recalculation(self):
        original = saved_budget_formulas({'预算明细!D2': "='预算总览'!$B$5*C2",
                                          '预算总览!B8': "=SUM('预算明细'!D2:D2)"})
        saved = saved_budget_formulas({'预算明细!D2': '=预算总览!$B$5*C2',
                                       '预算总览!B8': '=SUM(预算明细!D2:D2)'})
        result = check_xlsx_result(xlsx_cells(original), saved)
        self.assertEqual(result['formula_count'], 2)
        self.assertTrue(result['formula_caches_checked'])
        # The reader accepts the same local reference with quotes added too.
        self.assertEqual(check_xlsx_result(xlsx_cells(saved), original), result)

    def test_sheet_quote_exception_rejects_other_formula_changes(self):
        original = saved_budget_formulas({'预算明细!D2': "='预算总览'!$B$5*C2"})
        before = xlsx_cells(original)
        for formula in ('=预算总览!$B$6*C2', '=预算总览!B5*C2', '=预算总览!$B$5+C2',
                        '=预算总览!$B$5*C3', '=预算明细!$B$5*C2', '=6'):
            with self.subTest(formula=formula), self.assertRaisesRegex(ValueError, 'formulas changed'):
                check_xlsx_result(before, saved_budget_formulas({'预算明细!D2': formula}))
        for formulas in ({}, {'预算明细!D2': '6'},
                         {'预算明细!D2': '=预算总览!$B$5*C2', '预算明细!D3': '=1'}):
            with self.subTest(formulas=formulas), self.assertRaisesRegex(ValueError, 'formulas changed'):
                check_xlsx_result(before, saved_budget_formulas(formulas))
        with self.assertRaisesRegex(ValueError, 'uncalculated'):
            check_xlsx_result(before, saved_budget_formulas({'预算明细!D2': '=预算总览!$B$5*C2'}, cached=None))

    def test_sheet_quote_exception_does_not_rewrite_strings_or_unsupported_references(self):
        for original, saved, sheets in (
            ('="\'预算总览\'!B5"', '="预算总览!B5"', ()),
            ('=INDIRECT("\'预算总览\'!B5")', '=INDIRECT("预算总览!B5")', ()),
            ("='不存在'!B5", '=不存在!B5', ()),
            ("='预算 总览'!B5", '=预算 总览!B5', ('预算 总览',)),
            ("='预算''总览'!B5", "=预算'总览!B5", ("预算'总览",)),
            ("='[other.xlsx]预算总览'!B5", '=[other.xlsx]预算总览!B5', ()),
            ("='预算总览:预算明细'!B5", '=预算总览:预算明细!B5', ()),
            ("='预算总览'!Rate", '=预算总览!Rate', ()),
            ("=SUM(  '预算总览'!B5)", '=SUM( 预算总览!B5)', ()),
        ):
            with self.subTest(original=original):
                before = xlsx_cells(saved_budget_formulas({'预算明细!D2': original}, extra_sheets=sheets))
                with self.assertRaisesRegex(ValueError, 'formulas changed'):
                    check_xlsx_result(before, saved_budget_formulas({'预算明细!D2': saved}, extra_sheets=sheets))

    def test_saved_xlsx_requires_formula_type_and_cache_but_keeps_equals_text(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'fictional.xlsx'
            fictional_budget(source)
            data = source.read_bytes()
            before = xlsx_cells(data)
            with self.assertRaisesRegex(ValueError, 'uncalculated'):
                check_xlsx_result(before, data)

            def with_caches(*, text_formula=False, error=False):
                result = io.BytesIO()
                ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(result, 'w') as target:
                    for name in original.namelist():
                        value = original.read(name)
                        if name == 'xl/worksheets/sheet1.xml':
                            xml = ET.fromstring(value)
                            for cell in xml.iter(ns + 'c'):
                                formula = cell.find(ns + 'f')
                                if formula is not None:
                                    cell.find(ns + 'v').text = '0'
                                    if error and cell.attrib['r'] == 'D2':
                                        cell.set('t', 'e')
                                        cell.find(ns + 'v').text = '#DIV/0!'
                                    if text_formula and cell.attrib['r'] == 'D2':
                                        cell.remove(formula)
                                        cell.set('t', 'str')
                                        cell.find(ns + 'v').text = '=B2*C2'
                            value = ET.tostring(xml)
                        target.writestr(name, value)
                return result.getvalue()

            result = check_xlsx_result(before, with_caches())
            self.assertEqual(result['formula_count'], 7)
            self.assertEqual(result['equals_text_cells'], ['虚构婚宴预算!B9', '虚构婚宴预算!B10'])
            with self.assertRaisesRegex(ValueError, 'formulas changed'):
                check_xlsx_result(before, with_caches(text_formula=True))
            with self.assertRaisesRegex(ValueError, 'error'):
                check_xlsx_result(before, with_caches(error=True))

    def test_xlsx_source_is_supported_without_changing_saved_formulas(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional.xlsx'
            fictional_budget(source)
            original = source.read_bytes()
            formulas = load_workbook(source, data_only=False)
            values = load_workbook(source, data_only=True)
            self.assertEqual(formulas.active['D2'].data_type, 'f')
            self.assertIsNone(values.active['D2'].value, 'fixture must reproduce a missing formula cache')
            self.assertEqual(formulas.active['B9'].data_type, 's')
            self.assertEqual(formulas.active['B9'].value, '=SUM(D2:D5)')
            self.assertEqual(prepare_paths(root, source, root / 'outputs' / 'calculated')[0], source.resolve())
            self.assertEqual(source.read_bytes(), original)

    def test_existing_render_directory_reports_safe_recovery_without_conversion(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional.pptx'
            source.write_bytes(b'original')
            target = root / 'outputs' / 'existing'
            target.mkdir()
            for populated in (False, True):
                if populated:
                    (target / 'previous.pdf').write_bytes(b'previous result')
                error = io.StringIO()
                args = ['render', '--work-root', str(root), '--input', str(source), '--output-dir', str(target)]
                with patch('sys.argv', args), patch('team_file_render.os') as identity, \
                        patch('team_file_render.run') as convert, redirect_stderr(error):
                    identity.name = 'posix'
                    identity.geteuid.return_value = 1001
                    with self.assertRaises(SystemExit) as stopped:
                        main()
                self.assertEqual(stopped.exception.code, 1)
                self.assertIn('Output directory already exists', error.getvalue())
                self.assertIn('do not create it beforehand', error.getvalue())
                self.assertNotIn(str(root), error.getvalue())
                convert.assert_not_called()
                self.assertEqual(source.read_bytes(), b'original')
                if populated:
                    self.assertEqual((target / 'previous.pdf').read_bytes(), b'previous result')

    def test_paths_preserve_originals_and_reject_wrong_work_or_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional.docx'
            source.write_bytes(b'fictional')
            self.assertEqual(prepare_paths(root, source, root / 'outputs' / 'v1')[0], source.resolve())
            for target in (root, root / 'outputs', root / 'elsewhere'):
                with self.assertRaises(ValueError):
                    prepare_paths(root, source, target)
            wrong = root / 'unselected.docx'
            wrong.write_bytes(b'fictional')
            with self.assertRaises(ValueError):
                prepare_paths(root, wrong, root / 'outputs' / 'v1')

    def test_office_preflight_requires_exact_protocol(self):
        with patch('team_file_render.run', return_value=b'wrong'):
            with self.assertRaises(ValueError):
                office_check('/candidate/gateway', 'fictional.docx', b'fictional')

    @unittest.skipUnless(os.name == 'posix', 'process-group cleanup is Linux runtime behavior')
    def test_timeout_kills_transient_process_group(self):
        with tempfile.TemporaryDirectory() as temp:
            pid_file = Path(temp) / 'child.pid'
            started = time.monotonic()
            with self.assertRaises(subprocess.TimeoutExpired):
                run(['/bin/sh', '-c', 'sleep 30 & echo $! > "$1"; wait', 'fixture', str(pid_file)], timeout=.5)
            self.assertLess(time.monotonic() - started, 3)
            status = Path('/proc') / pid_file.read_text().strip() / 'stat'
            for _ in range(100):
                try:
                    state = status.read_text().split(') ')[1]
                except (FileNotFoundError, ProcessLookupError):
                    break  # Reaping can race opening/reading the proc entry.
                if state.startswith('Z'):
                    break
                time.sleep(.01)
            else:
                self.fail('spawned child remains alive after timeout')

    @unittest.skipUnless(os.name == 'posix' and os.geteuid() != 0, 'ordinary Linux worker limits')
    def test_poppler_inherits_limits_before_exec(self):
        probe = '''
import json, resource
from unittest.mock import patch
from team_file_render import preview_worker
with patch('team_file_render.os.execv') as execute:
    preview_worker('/fixture/input.pdf', '/fixture/page')
    assert execute.call_args.args[0] == '/usr/bin/pdftoppm'
    print(json.dumps([resource.getrlimit(key) for key in (resource.RLIMIT_AS, resource.RLIMIT_CPU, resource.RLIMIT_FSIZE, resource.RLIMIT_CORE)]))
'''
        output = subprocess.check_output([sys.executable, '-B', '-c', probe], cwd=Path(__file__).parent)
        self.assertEqual(json.loads(output), [[2 << 30] * 2, [90] * 2, [128 << 20] * 2, [0, 0]])


@unittest.skipUnless(os.environ.get('TEAM_FILES_TEST_XLSX_RECALC') == '1'
                     and os.environ.get('TEAM_FILES_TEST_GATEWAY') and sys.platform in ('linux', 'darwin'),
                     'requires explicit ordinary Linux/Mac Office fixture and matching gateway')
class NativeXlsxRenderTests(unittest.TestCase):
    def test_optional_sheet_quotes_recalculate_and_preview(self):
        self.assertNotEqual(os.geteuid(), 0)
        if sys.platform == 'linux':
            deny_unix_sockets()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp).resolve()
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional-budget.xlsx'
            book = Workbook()
            book.active.title = '预算总览'
            book.create_sheet('预算明细')
            book['预算总览']['B5'] = 2
            book['预算明细']['C2'] = 3
            book['预算明细']['D2'] = "='预算总览'!$B$5*C2"
            book['预算总览']['B8'] = "=SUM('预算明细'!D2:D2)"
            book.save(source)
            original = source.read_bytes()
            target = root / 'outputs' / 'recalculated'
            result = render(root, source, target, os.environ['TEAM_FILES_TEST_GATEWAY'])
            saved = target / source.name
            formulas = load_workbook(saved, data_only=False)
            values = load_workbook(saved, data_only=True)
            try:
                self.assertIn(formulas['预算明细']['D2'].value, ("='预算总览'!$B$5*C2", '=预算总览!$B$5*C2'))
                self.assertIn(formulas['预算总览']['B8'].value, ("=SUM('预算明细'!D2:D2)", '=SUM(预算明细!D2:D2)'))
                for sheet, cell in (('预算明细', 'D2'), ('预算总览', 'B8')):
                    self.assertEqual(formulas[sheet][cell].data_type, 'f')
                    self.assertEqual(values[sheet][cell].value, 6)
            finally:
                formulas.close(); values.close()
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(result['xlsx']['formula_count'], 2)
            self.assertTrue(result['xlsx']['formula_caches_checked'])
            pages = len(PdfReader(saved.with_suffix('.pdf')).pages)
            self.assertGreater(pages, 0)
            self.assertEqual(len(list(target.glob('page-*.png'))), pages)
            self.assertEqual(json.loads((target / 'render.json').read_text(encoding='utf-8'))['pages'], pages)


@unittest.skipUnless(os.environ.get('TEAM_FILES_TEST_GATEWAY'), 'requires matching candidate binary and Linux Office packages')
class ActualRenderTests(unittest.TestCase):
    def test_xlsx_recalculation_preserves_formula_types_text_and_original(self):
        self.assertNotEqual(os.geteuid(), 0)
        deny_unix_sockets()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional-budget.xlsx'
            fictional_budget(source)
            original = source.read_bytes()
            result = render(root, source, root / 'outputs' / 'v1', os.environ['TEAM_FILES_TEST_GATEWAY'])
            saved = root / 'outputs' / 'v1' / source.name
            formulas = load_workbook(saved, data_only=False)
            values = load_workbook(saved, data_only=True)
            for address, formula, expected in (
                ('D2', '=B2*C2', 89760), ('D3', '=B3*C3', 3960),
                ('D6', '=SUM(D2:D5)', 103720), ('D7', '=110000-D6', 6280),
                ('D8', '=IF(D6<=110000,"预算内","超预算")', '预算内'),
                ('D11', "='依据 v2'!B1", 4080), ('D12', '=IF(D6>0,"","")', None),
            ):
                self.assertEqual(formulas.active[address].data_type, 'f')
                self.assertEqual(formulas.active[address].value, formula)
                self.assertEqual(values.active[address].value, expected)
            for address in ('B9', 'B10'):
                self.assertEqual(formulas.active[address].data_type, 's')
            self.assertEqual(formulas.active['B9'].value, '=SUM(D2:D5)')
            self.assertEqual(formulas.active['B10'].value, '= 表示预算口径说明，不是公式')
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(result['xlsx']['formula_count'], 7)
            self.assertTrue(result['xlsx']['formula_caches_checked'])
            self.assertEqual(len(result['xlsx']['equals_text_cells']), 2)
            self.assertTrue(result['xlsx']['equals_text_requires_task_review'])
            self.assertEqual(result['status'], 'rendered_not_visually_verified')
            text = ''.join(page.extract_text() for page in PdfReader(saved.with_suffix('.pdf')).pages)
            self.assertIn('虚构婚宴预算', text)
            self.assertIn('103,720', text)
            self.assertEqual(len(list(saved.parent.glob('page-*.png'))), result['pages'])

    def test_xlsx_formula_errors_stop_before_success_manifest(self):
        self.assertNotEqual(os.geteuid(), 0)
        deny_unix_sockets()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'fictional-error.xlsx'
            workbook = Workbook()
            workbook.active['A1'] = '=1/0'
            workbook.save(source)
            original = source.read_bytes()
            with self.assertRaises(ValueError):
                render(root, source, root / 'outputs' / 'rejected', os.environ['TEAM_FILES_TEST_GATEWAY'])
            self.assertFalse((root / 'outputs' / 'rejected' / 'render.json').exists())
            self.assertEqual(source.read_bytes(), original)

    def test_xlsx_active_content_is_rejected_before_office_starts(self):
        self.assertNotEqual(os.geteuid(), 0)
        deny_unix_sockets()
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            for content in ('external-link', 'macro'):
                with self.subTest(content=content):
                    source = root / 'inputs' / f'{content}.xlsx'
                    workbook = Workbook()
                    workbook.active['A1'] = '虚构说明'
                    if content == 'external-link':
                        workbook.active['A1'].hyperlink = 'https://example.invalid/never-contact'
                    workbook.save(source)
                    if content == 'macro':
                        with zipfile.ZipFile(source, 'a') as package:
                            package.writestr('xl/vbaProject.bin', b'fictional rejected macro marker')
                    original = source.read_bytes()
                    with patch('team_file_render.run', wraps=run) as tools:
                        with self.assertRaises(RuntimeError):
                            render(root, source, root / 'outputs' / content, os.environ['TEAM_FILES_TEST_GATEWAY'])
                    self.assertEqual(tools.call_count, 1)
                    self.assertEqual(tools.call_args.args[0][1], 'file-check')
                    self.assertFalse((root / 'outputs' / content).exists())
                    self.assertEqual(source.read_bytes(), original)

    def test_chinese_slides_and_word_toc_revision(self):
        self.assertNotEqual(os.geteuid(), 0, 'acceptance must use an ordinary identity')
        deny_unix_sockets()
        gateway = os.environ['TEAM_FILES_TEST_GATEWAY']
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            pages = []
            for version in (1, 2):
                source = root / 'inputs' / f'fictional-v{version}.docx'
                fictional_template(source, extra=version == 2)
                original = hashlib.sha256(source.read_bytes()).hexdigest()
                target = root / 'outputs' / f'v{version}'
                result = render(root, source, target, gateway)
                self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)
                reader = PdfReader(target / (source.stem + '.pdf'))
                text = [p.extract_text() for p in reader.pages]
                self.assertIn('第一章', text[0])
                self.assertIn('第二章', text[0])
                self.assertNotIn('Error!', '\n'.join(text))
                self.assertEqual(result['pages'], len(text))
                self.assertTrue((target / (source.stem + '.docx')).is_file())
                with zipfile.ZipFile(target / (source.stem + '.docx')) as package:
                    xml = ET.fromstring(package.read('word/document.xml'))
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    saved_paragraphs = [''.join(p.itertext()) for p in xml.findall('.//w:p', ns)]
                for heading in ('第一章 客户需求', '第二章 执行安排'):
                    actual = next(number for number, page in enumerate(text[1:], 2) if heading in page)
                    entry = re.search(re.escape(heading) + r'\.+\s*(\d+)', text[0])
                    self.assertIsNotNone(entry, 'rendered TOC entry missing')
                    self.assertEqual(int(entry.group(1)), actual, 'TOC points to wrong PDF page')
                    self.assertTrue(any(p.endswith(heading + str(actual)) for p in saved_paragraphs),
                                    'saved DOCX TOC differs from actual page')
                for number, page in enumerate(text, 1):
                    self.assertIn(f'Page {number} of {len(text)}', ' '.join(page.split()))
                pages.append(len(text))
            self.assertGreater(pages[1], pages[0])
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
            shape.text = '虚构婚宴方案：22 桌，预算 101,520 元'
            for paragraph in shape.text_frame.paragraphs:
                for item in paragraph.runs:
                    item.font.name = 'Noto Sans CJK SC'
            source = root / 'inputs' / 'fictional-slides.pptx'
            deck.save(source)
            result = render(root, source, root / 'outputs' / 'slides', gateway)
            self.assertEqual(result['pages'], 1)
            text = PdfReader(root / 'outputs' / 'slides' / 'fictional-slides.pdf').pages[0].extract_text()
            self.assertIn('虚构婚宴', text)
            self.assertIn('101,520', text)
            self.assertEqual(Presentation(source).slides[0].shapes[0].text, shape.text)


@unittest.skipUnless(os.name == 'posix' and os.geteuid() == 0 and os.environ.get('TEAM_FILES_TEST_GATEWAY'),
                     'requires disposable Linux numeric-UID initializer and document packages')
class HandoffRenderTests(unittest.TestCase):
    def test_model_render_and_revision_readable_by_ordinary_gateway_only(self):
        # Root only provisions the fixture; rendering and reading drop all groups
        # except the same restricted handoff group used by the deployed accounts.
        gateway_uid, model_uid, shared_gid = 61000, 61001, 61002
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            root.chmod(0o755)
            work = root / 'work'
            work.mkdir()
            os.chown(work, gateway_uid, shared_gid)
            work.chmod(0o2750)
            for name, mode in (('inputs', 0o2750), ('outputs', 0o2770)):
                path = work / name
                path.mkdir()
                os.chown(path, gateway_uid, shared_gid)
                path.chmod(mode)

            def as_user(uid, script, *args, group=shared_gid):
                return subprocess.run([sys.executable, '-B', '-c', script, str(work), *args],
                                      user=uid, group=group, extra_groups=[], umask=0o077,
                                      cwd=Path(__file__).parent, capture_output=True, text=True, timeout=120)

            made = as_user(model_uid, '''
import json, os, sys
from pathlib import Path
from test_team_file_render import fictional_template, fictional_budget, deny_unix_sockets
from team_file_render import render
assert os.geteuid() == 61001
assert 'CapEff:\t0000000000000000' in Path('/proc/self/status').read_text()
deny_unix_sockets()
root = Path(sys.argv[1])
legacy = root / 'outputs' / 'legacy-private'
legacy.mkdir(mode=0o700)
(legacy / 'result').write_text('fictional')
for version in (1, 2):
    source = root / 'outputs' / f'source-v{version}.docx'
    fictional_template(source, extra=version == 2)
    render(root, source, root / 'outputs' / f'v{version}', os.environ['TEAM_FILES_TEST_GATEWAY'])
source = root / 'outputs' / 'budget.xlsx'
fictional_budget(source)
render(root, source, root / 'outputs' / 'budget-v1', os.environ['TEAM_FILES_TEST_GATEWAY'])
''')
            self.assertEqual(made.returncode, 0, made.stderr)
            checked = as_user(gateway_uid, '''
import hashlib, json, os, sys
from pathlib import Path
assert os.geteuid() == 61000
assert 'CapEff:\t0000000000000000' in Path('/proc/self/status').read_text()
root = Path(sys.argv[1]) / 'outputs'
try:
    (root / 'legacy-private' / 'result').read_bytes()
except PermissionError:
    pass
else:
    raise AssertionError('ordinary host bypassed old private output')
pages = []
for version in (1, 2):
    target = root / f'v{version}'
    record = json.loads((target / 'render.json').read_text())
    for name, digest in record['files'].items():
        path = target / name
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
        assert path.stat().st_uid == 61001 and path.stat().st_gid == 61002
        assert path.stat().st_mode & 0o777 == 0o640
    pages.append(record['pages'])
assert pages[1] > pages[0]
budget = root / 'budget-v1'
record = json.loads((budget / 'render.json').read_text())
assert record['xlsx']['formula_count'] == 7
for name, digest in record['files'].items():
    path = budget / name
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    assert path.stat().st_uid == 61001 and path.stat().st_gid == 61002
    assert path.stat().st_mode & 0o777 == 0o640
print(json.dumps(pages))
''')
            self.assertEqual(checked.returncode, 0, checked.stderr)
            self.assertGreater(json.loads(checked.stdout)[1], json.loads(checked.stdout)[0])
            denied = as_user(61003, 'from pathlib import Path; import sys; (Path(sys.argv[1]) / "outputs/v1/render.json").read_bytes()', group=61003)
            self.assertNotEqual(denied.returncode, 0)
            self.assertIn('PermissionError', denied.stderr)


if __name__ == '__main__':
    unittest.main()
