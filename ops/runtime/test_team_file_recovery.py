"""Fictional work files at the context/render CLI boundary; no real delivery."""
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from contextlib import redirect_stdout, redirect_stderr
from unittest.mock import patch
from openpyxl import Workbook, load_workbook

import team_file_render as renderer


class ContextSourceTests(unittest.TestCase):
    def test_relative_reference_cannot_escape_work(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            (root / 'outputs').mkdir()
            (root / 'outside.docx').write_bytes(b'private')
            context = {'output_dir': str(root / 'outputs')}
            for value in ('outputs/../outside.docx', 'outputs//outside.docx',
                          'outputs/../../outside.docx', 'outputs/..\\outside.docx', str(root / 'outside.docx')):
                with self.subTest(value=value), self.assertRaises(ValueError):
                    renderer.work_path(context, value)

    @unittest.skipUnless(os.name == 'posix', 'POSIX link boundary')
    def test_linked_output_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'outputs').mkdir()
            (root / 'private.docx').write_bytes(b'private')
            (root / 'outputs' / 'link.docx').symlink_to(root / 'private.docx')
            with self.assertRaises(ValueError):
                renderer.work_path({'output_dir': str(root / 'outputs')}, 'outputs/link.docx')

    def test_saved_revision_resolves_without_copying_work_root(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'outputs' / 'final.docx'
            source.write_bytes(b'fictional saved revision')
            context = {'enabled': True, 'work_id': 'current-work',
                       'output_dir': str(root / 'outputs'), 'inputs': [],
                       'latest_version': 0, 'artifacts': []}
            reply = subprocess.CompletedProcess([], 0, json.dumps(context).encode(), b'')
            with patch('team_file_render.subprocess.run', return_value=reply):
                work, selected, target, binding = renderer.context_source(None, 'final.docx')
            self.assertEqual(selected, source)
            self.assertEqual(work, root)
            self.assertEqual(target.parent, root / 'outputs')
            self.assertFalse(target.exists())
            self.assertEqual(binding['source_path'], 'outputs/final.docx')
            self.assertEqual(binding['work_id'], 'current-work')
            self.assertEqual(binding['source_sha256'], hashlib.sha256(source.read_bytes()).hexdigest())
            self.assertEqual(renderer.prepare_paths(work, selected, target)[0], source)

    def test_source_id_is_exact_and_preserved_not_manually_rebuilt(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            (root / 'inputs').mkdir()
            (root / 'outputs').mkdir()
            source = root / 'inputs' / 'original.pdf'
            source.write_bytes(b'fictional selected original')
            item = {'id': 'a' * 63 + 'c', 'name': 'original.pdf', 'path': str(source),
                    'sha256': hashlib.sha256(source.read_bytes()).hexdigest()}
            context = {'enabled': True, 'work_id': 'current-work',
                       'output_dir': str(root / 'outputs'), 'inputs': [item]}
            reply = subprocess.CompletedProcess([], 0, json.dumps(context).encode(), b'')
            with patch('team_file_render.subprocess.run', return_value=reply):
                _, selected, _, binding = renderer.context_source(item['id'], None)
                self.assertEqual(selected, source)
                self.assertEqual(binding['source_input'], {k: item[k] for k in ('id', 'name', 'sha256')})
                with self.assertRaises(ValueError):
                    renderer.context_source(item['id'][:-1], None)
                source.write_bytes(b'changed original')
                with self.assertRaises(ValueError):
                    renderer.context_source(item['id'], None)


class PreparedDeliveryTests(unittest.TestCase):
    def fixture(self, root):
        (root / 'inputs').mkdir()
        output = root / 'outputs'
        target = output / 'render-1'
        target.mkdir(parents=True)
        (output / 'revision.docx').write_bytes(b'last saved revision')
        (target / 'revision.docx').write_bytes(b'updated fields')
        (target / 'revision.pdf').write_bytes(b'final pages')
        (target / 'page-1.png').write_bytes(b'first page pixels')
        manifest = {'status': 'rendered_not_visually_verified', 'work_id': 'this-work',
                    'source_path': 'outputs/revision.docx',
                    'source_sha256': hashlib.sha256((output / 'revision.docx').read_bytes()).hexdigest(),
                    'pages': 1,
                    'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in target.iterdir()}}
        (target / 'render.json').write_text(json.dumps(manifest), encoding='utf-8')
        context = {'enabled': True, 'work_id': 'this-work', 'output_dir': str(output),
                   'latest_version': 3, 'inputs': [], 'artifacts': []}
        return context, target

    def test_revision_cannot_reuse_old_render_or_incomplete_inspection(self):
        with tempfile.TemporaryDirectory() as tmp:
            context, target = self.fixture(Path(tmp).resolve())
            with patch('team_file_render.subprocess.run', return_value=subprocess.CompletedProcess(
                    [], 0, json.dumps(context).encode(), b'')):
                with self.assertRaises(ValueError):
                    renderer.prepare_delivery('render-1/render.json', 'revision.docx', [], ['revision.docx', 'revision.pdf'])
                receipt = renderer.prepare_delivery('render-1/render.json', 'revision.docx',
                                                    ['page-1.png'], ['revision.docx', 'revision.pdf'])
                request = json.loads(Path(receipt['request_file']).read_text(encoding='utf-8'))
                self.assertEqual(request, {'work_id': 'this-work', 'expected_version': 3,
                                           'path': 'render-1/revision.docx',
                                           'sha256': hashlib.sha256(b'updated fields').hexdigest()})
                (target.parent / 'revision.docx').write_bytes(b'new layout revision')
                with self.assertRaises(ValueError):
                    renderer.prepare_delivery('render-1/render.json', 'revision.docx',
                                              ['page-1.png'], ['revision.docx', 'revision.pdf'])

    def test_unknown_delivery_prepares_status_read_not_a_send(self):
        with tempfile.TemporaryDirectory() as tmp:
            context, _ = self.fixture(Path(tmp).resolve())
            context['artifacts'] = [{'delivery_id': 'uncertain-delivery', 'status': 'unknown',
                                     'path': 'other.pdf', 'sha256': 'b' * 64}]
            calls = []
            def host(args, **kwargs):
                calls.append(args[-1])
                value = context if args[-1] == 'work-context' else {'status': 'unknown', 'delivery_id': 'uncertain-delivery'}
                return subprocess.CompletedProcess(args, 0, json.dumps(value).encode(), b'')
            with patch('team_file_render.subprocess.run', side_effect=host):
                result = renderer.prepare_delivery('render-1/render.json', 'revision.docx',
                                                   ['page-1.png'], ['revision.docx', 'revision.pdf'])
            self.assertEqual(calls, ['work-context', 'file-status'])
            self.assertEqual(result['status'], 'reconcile_delivery')
            self.assertNotIn('request_file', result)

    def test_changed_outputs_wrong_work_and_missing_pages_do_not_prepare(self):
        for defect in ('output', 'work', 'pages', 'manifest_type'):
            with self.subTest(defect=defect), tempfile.TemporaryDirectory() as tmp:
                context, target = self.fixture(Path(tmp).resolve())
                manifest = json.loads((target / 'render.json').read_bytes())
                if defect == 'output':
                    (target / 'revision.pdf').write_bytes(b'edited after render')
                elif defect == 'work':
                    manifest['work_id'] = 'other-work'
                elif defect == 'pages':
                    manifest['pages'] = 2
                else:
                    manifest = []
                (target / 'render.json').write_text(json.dumps(manifest), encoding='utf-8')
                with patch('team_file_render.subprocess.run', return_value=subprocess.CompletedProcess(
                        [], 0, json.dumps(context).encode(), b'')), self.assertRaises(ValueError):
                    renderer.prepare_delivery('render-1/render.json', 'revision.docx',
                                              ['page-1.png'], ['revision.docx', 'revision.pdf'])
                self.assertEqual(list(target.glob('delivery-request-*')), [])

    def test_public_cli_prepares_current_request_and_never_calls_send(self):
        with tempfile.TemporaryDirectory() as tmp:
            context, _ = self.fixture(Path(tmp).resolve())
            calls = []
            def host(args, **kwargs):
                calls.append(args[-1])
                return subprocess.CompletedProcess(args, 0, json.dumps(context).encode(), b'')
            stdout = io.StringIO()
            with patch('sys.argv', ['render', '--prepare-delivery', 'render-1/render.json',
                                   '--artifact', 'revision.docx', '--inspected-page', 'page-1.png',
                                   '--read-content', 'revision.docx', '--read-content', 'revision.pdf']), \
                    patch('team_file_render.subprocess.run', side_effect=host), redirect_stdout(stdout):
                renderer.main()
            self.assertEqual(calls, ['work-context'])
            self.assertEqual(json.loads(stdout.getvalue())['status'], 'prepared_not_sent')

    def test_permission_denial_and_missing_client_never_become_source_retry(self):
        for error in (PermissionError('private diagnostic'), FileNotFoundError('missing client')):
            with self.subTest(error=type(error).__name__), \
                    patch('sys.argv', ['render', '--output-file', 'final.docx']), \
                    patch('team_file_render.subprocess.run', side_effect=error), \
                    redirect_stderr(io.StringIO()) as stderr, self.assertRaises(SystemExit):
                try:
                    renderer.main()
                finally:
                    self.assertNotIn('input_missing_refresh_work_context', stderr.getvalue())
                    self.assertNotIn(str(error), stderr.getvalue())


class CellIntentTests(unittest.TestCase):
    def test_explanatory_formula_string_is_rejected_until_saved_as_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / 'revision.xlsx'
            book = Workbook()
            book.active.title = '来源'
            book.active['A1'] = '=SUM(F6:F41)，缓存50,400'
            book.active['B1'] = '=2*3'
            book.save(p)
            before = renderer.xlsx_cells(p.read_bytes())
            with self.assertRaises(ValueError):
                renderer.check_text_cells(before, ['来源!A1'])
            book.active['A1'].data_type = 's'
            book.save(p)
            renderer.check_text_cells(renderer.xlsx_cells(p.read_bytes()), ['来源!A1'])
            self.assertTrue(renderer.xlsx_cells(p.read_bytes())['来源!B1']['is_formula'])
            with self.assertRaises(ValueError):
                renderer.check_xlsx_result(before, p.read_bytes())


class SourceLedgerTests(unittest.TestCase):
    fixture = PreparedDeliveryTests.fixture
    def test_page_rows_keep_machine_source_identity_and_do_not_expand_to_other_pages(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            context, target = self.fixture(root)
            original = root / 'inputs' / 'original.pdf'
            original.write_bytes(b'fictional two page original')
            item = {'id': 'a' * 63 + 'c', 'name': 'original.pdf', 'path': str(original),
                    'sha256': hashlib.sha256(original.read_bytes()).hexdigest()}
            context['inputs'] = [item]
            m = json.loads((target / 'render.json').read_bytes())
            m.update(source_path='inputs/original.pdf', source_sha256=item['sha256'],
                     source_input={k: item[k] for k in ('id', 'name', 'sha256')})
            regions = []
            for number in range(1, 4):
                name = f'detail-1-{number}.png'
                (target / name).write_bytes(b'detail pixels')
                m['files'][name] = hashlib.sha256(b'detail pixels').hexdigest()
                regions.append({'file': name, 'sha256': m['files'][name]})
            m['source_details'] = [{'page': 1, 'page_file': 'page-1.png',
                                    'page_sha256': m['files']['page-1.png'], 'regions': regions}]
            (target / 'render.json').write_text(json.dumps(m), encoding='utf-8')
            rows = [{'page': 1, 'quote': '本页未附银行凭据', 'meaning': '仅指本页依据，不推断其他页'}]
            (target.parent / 'quotes.json').write_text(json.dumps(rows), encoding='utf-8')
            with patch('team_file_render.subprocess.run', return_value=subprocess.CompletedProcess(
                    [], 0, json.dumps(context).encode(), b'')):
                output = renderer.source_ledger('render-1/render.json', 'quotes.json')
                saved = json.loads(Path(output['ledger']).read_text(encoding='utf-8'))
                self.assertEqual(saved['source'], {k: item[k] for k in ('id', 'name', 'sha256')})
                self.assertEqual(saved['rows'][0]['page'], 1)
                self.assertEqual(saved['rows'][0]['quote'], '本页未附银行凭据')
                self.assertIn(item['id'], saved['rows'][0]['citation'])
                rows[0]['page'] = 2
                (target.parent / 'quotes.json').write_text(json.dumps(rows), encoding='utf-8')
                with self.assertRaises(ValueError):
                    renderer.source_ledger('render-1/render.json', 'quotes.json')


@unittest.skipUnless(os.name == 'posix' and getattr(os, 'geteuid', lambda: 0)() != 0
                     and os.environ.get('TEAM_FILES_TEST_GATEWAY'),
                     'requires explicit ordinary native Office fixture and paired validator')
class NativeRecoveryTests(unittest.TestCase):
    def test_explanation_correction_recalculates_and_revision_needs_new_render(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp).resolve()
            (root / 'inputs').mkdir()
            output = root / 'outputs'
            output.mkdir()
            context = {'enabled': True, 'work_id': 'native-fixture', 'output_dir': str(output),
                       'inputs': [], 'artifacts': [], 'latest_version': 0}
            book = Workbook()
            book.active.title = 'Source'
            book.active['A1'] = '=2*3'
            book.active['B1'] = '=2*3'
            book.save(output / 'r1.xlsx')
            def invoke(name, failure=None):
                stdout = io.StringIO()
                with patch('sys.argv', ['render', '--output-file', name, '--text-cell', 'Source!A1',
                                       '--gateway', os.environ['TEAM_FILES_TEST_GATEWAY']]), \
                        patch('team_file_render.subprocess.run', return_value=subprocess.CompletedProcess(
                            [], 0, json.dumps(context).encode(), b'')), redirect_stdout(stdout):
                    if failure is None:
                        renderer.main()
                    else:
                        with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                            renderer.main()
                        failure.update(json.loads(stdout.getvalue()))
                        return
                return json.loads(stdout.getvalue())
            with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                invoke('r1.xlsx')
            self.assertEqual(list(output.glob('render-*')), [])
            book.active['A1'].data_type = 's'
            book.active['B1'] = '=1/0'
            book.save(output / 'bad-calculation.xlsx')
            failed = {}
            invoke('bad-calculation.xlsx', failed)
            self.assertEqual(failed['code'], 'xlsx_result_invalid')
            retained = Path(failed['retained_workbook'])
            self.assertTrue(retained.is_relative_to(output))
            self.assertEqual(failed['retained_sha256'], hashlib.sha256(retained.read_bytes()).hexdigest())
            self.assertFalse((retained.parent / 'render.json').exists())
            self.assertFalse(failed['delivery_attempted'])
            self.assertEqual(renderer.xlsx_cells(retained.read_bytes())['Source!B1']['type'], 'e')
            book.active['B1'] = '=2*3'
            book.save(output / 'r2.xlsx')
            receipt = invoke('r2.xlsx')
            target = Path(receipt['manifest']).parent
            with (target / 'r2.xlsx').open('rb') as f:
                calculated = load_workbook(f, data_only=True)
                self.assertEqual(calculated['Source']['B1'].value, 6)
                self.assertEqual(calculated['Source']['A1'].value, '=2*3')
                calculated.close()
            book.active['B1'] = '=3*3'
            book.save(output / 'r2.xlsx')
            with patch('team_file_render.subprocess.run', return_value=subprocess.CompletedProcess(
                    [], 0, json.dumps(context).encode(), b'')), self.assertRaises(ValueError):
                renderer.prepare_delivery(receipt['manifest_ref'], 'r2.xlsx', ['page-1.png'], ['r2.xlsx', 'r2.pdf'])
            book.save(output / 'r3.xlsx')
            revised = invoke('r3.xlsx')
            with (Path(revised['manifest']).parent / 'r3.xlsx').open('rb') as f:
                calculated = load_workbook(f, data_only=True)
                self.assertEqual(calculated['Source']['B1'].value, 9)
                calculated.close()
            self.assertTrue((output / 'r1.xlsx').is_file())
            self.assertTrue((target / 'r2.xlsx').is_file())


if __name__ == '__main__':
    unittest.main()
