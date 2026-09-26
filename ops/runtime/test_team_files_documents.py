"""Fictional Office round-trip; no model, Feishu, or company service is invoked."""
import hashlib
from pathlib import Path
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook


class DocumentRoundTripTests(unittest.TestCase):
    def test_fictional_analysis_and_revision_preserve_originals(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'fictional-orders.xlsx'
            workbook = Workbook()
            workbook.active.append(['Item', 'Units', 'Unit price'])
            workbook.active.append(['Fictional Cedar', 3, 20])
            workbook.save(source)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            # This is a selected fictional source, not a live knowledge read.
            material = root / 'fictional-policy.txt'
            material.write_text('Fictional policy v1: summarize the total and state assumptions.', encoding='utf-8')
            source_text = material.read_text(encoding='utf-8')
            old_files = {}
            for version, units in ((1, 3), (2, 4)):
                report = load_workbook(source)
                report.active.title = f'Budget-v{version}'
                report.active['B2'] = units
                report.active['D1'] = 'Total'
                report.active['D2'] = '=B2*C2'
                report.active['E1'] = 'Explanation'
                report.active['E2'] = '=units multiplied by unit price'
                report.active['E2'].data_type = 's'
                xlsx = root / f'analysis-v{version}.xlsx'
                report.save(xlsx)
                doc = Document()
                doc.add_heading(f'Fictional analysis v{version}', 0)
                doc.add_paragraph(f'Total: {units * 20}. Source: fictional-orders.xlsx; fictional-policy.txt v1.')
                doc.add_paragraph(source_text)
                docx = root / f'analysis-v{version}.docx'
                doc.save(docx)
                saved = load_workbook(xlsx)
                self.assertEqual(saved.active.title, f'Budget-v{version}')
                self.assertEqual((saved.active['D2'].value, saved.active['D2'].data_type),
                                 ('=B2*C2', 'f'))
                self.assertEqual((saved.active['E2'].value, saved.active['E2'].data_type),
                                 ('=units multiplied by unit price', 's'))
                self.assertEqual(saved.active['B2'].value, units)
                saved.close()
                text = '\n'.join(p.text for p in Document(docx).paragraphs)
                self.assertIn(f'Fictional analysis v{version}', text)
                self.assertIn(f'Total: {units * 20}', text)
                self.assertIn('fictional-policy.txt v1', text)
                for path, digest in old_files.items():
                    self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), digest)
                old_files.update({p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (xlsx, docx)})
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)


if __name__ == '__main__':
    unittest.main()
