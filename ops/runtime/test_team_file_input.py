"""Read an actual registered workbook after a mistyped-path failure."""
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


SCRIPT = Path(__file__).resolve().parents[1] / 'skills/team-files/read_xlsx.py'


class RegisteredWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        (self.root / 'inputs').mkdir()
        (self.root / 'outputs').mkdir()
        self.path = self.root / 'inputs' / 'input_cca1c9dc20737faaedced451_quote.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = '报价'
        sheet.append(['数量', '单价', '合计'])
        sheet.append([3, 170, '=A2*B2'])
        book.create_sheet('说明').append(['预计，未确认'])
        book.save(self.path)
        self.raw = self.path.read_bytes()
        self.context = dict(enabled=True, work_id='fixture-work', output_dir=str(self.root / 'outputs'),
                            inputs=[dict(id='input-exact', name='quote.xlsx', path=str(self.path),
                                         sha256=hashlib.sha256(self.raw).hexdigest())])
        spec = importlib.util.spec_from_file_location('team_file_input', SCRIPT)
        self.reader = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.reader)
        self.calls = []

    def external(self, args, **kwargs):
        self.calls.append(args)
        if args == ['/usr/local/bin/cc-connect-file', 'work-context']:
            self.assertEqual(json.loads(kwargs['input']), {})
            return subprocess.CompletedProcess(args, 0, json.dumps(self.context).encode(), b'')
        self.assertEqual(args, ['/usr/local/bin/cc-connect-anc', 'file-check', 'source.xlsx'])
        self.assertEqual(kwargs['input'], self.raw)
        return subprocess.CompletedProcess(args, 0, b'OFFICE_OK_V1\n', b'')

    def run_reader(self, *extra):
        output = io.StringIO()
        with patch.object(self.reader.subprocess, 'run', side_effect=self.external), contextlib.redirect_stdout(output):
            code = self.reader.main(['--input-id', 'input-exact', *extra])
        return code, json.loads(output.getvalue())

    def test_mistyped_path_recovers_by_registered_id_and_reads_saved_cells(self):
        wrong = self.path.with_name(self.path.name.replace('aedced451', 'aedcd451'))
        with self.assertRaises(FileNotFoundError):
            load_workbook(wrong, read_only=True)
        code, result = self.run_reader('--sheet', '报价')
        self.assertEqual(code, 0)
        self.assertEqual(result['rows'][1]['cells']['A2'], 3)
        self.assertEqual(result['rows'][1]['cells']['C2'], {'formula': '=A2*B2', 'cached': None})
        self.assertEqual(result['worksheets'], ['报价', '说明'])
        self.assertIsNone(result['next_row'])
        self.assertEqual(result['source']['sha256'], hashlib.sha256(self.raw).hexdigest())
        self.assertEqual(self.path.read_bytes(), self.raw)
        self.assertEqual(len(self.calls), 2)

    def test_disabled_missing_id_and_hash_drift_are_not_read(self):
        for change, expected in [('disabled', 'file_work_disabled'), ('missing', 'input_id_not_unique'),
                                 ('hash', 'input_hash_changed')]:
            with self.subTest(change=change):
                before = json.loads(json.dumps(self.context))
                if change == 'disabled': self.context['enabled'] = False
                elif change == 'missing': self.context['inputs'][0]['id'] = 'other-id'
                else: self.context['inputs'][0]['sha256'] = '0' * 64
                self.calls.clear()
                code, result = self.run_reader()
                self.assertEqual(code, 1)
                self.assertEqual(result['error'], expected)
                self.assertNotIn('rows', result)
                self.assertEqual(len(self.calls), 1)
                self.context = before

    def test_input_outside_current_work_is_rejected_without_guessing(self):
        self.context['output_dir'] = str(self.root / 'different-work' / 'outputs')
        code, result = self.run_reader()
        self.assertEqual((code, result['error']), (1, 'input_outside_work'))
        self.assertEqual(len(self.calls), 1)

    def test_row_pages_have_explicit_continuation(self):
        code, first = self.run_reader('--sheet', '报价', '--max-rows', '1')
        self.assertEqual((code, first['next_row']), (0, 2))
        code, second = self.run_reader('--sheet', '报价', '--start-row', str(first['next_row']), '--max-rows', '1')
        self.assertEqual((code, second['next_row']), (0, None))
        self.assertEqual(second['rows'][0]['cells']['B2'], 170)

    def test_invalid_rows_missing_sheet_and_duplicate_id_fail_closed(self):
        for args, expected in [(('--start-row', '0'), 'invalid_row_range'),
                               (('--max-rows', '201'), 'invalid_row_range'),
                               (('--start-row', '3'), 'row_out_of_range'),
                               (('--sheet', 'absent'), 'sheet_not_found')]:
            with self.subTest(args=args):
                code, result = self.run_reader(*args)
                self.assertEqual((code, result['error']), (1, expected))
                self.assertNotIn('rows', result)
        self.context['inputs'].append(self.context['inputs'][0].copy())
        self.calls.clear()
        code, result = self.run_reader()
        self.assertEqual((code, result['error'], len(self.calls)), (1, 'input_id_not_unique', 1))

    def test_hardlinked_input_and_rejected_office_are_not_parsed(self):
        alias = self.path.with_name('alias.xlsx')
        os.link(self.path, alias)
        code, result = self.run_reader()
        self.assertEqual((code, result['error'], len(self.calls)), (1, 'unsafe_input_file', 1))
        alias.unlink()
        real_external = self.external
        def reject(args, **kwargs):
            reply = real_external(args, **kwargs)
            if args[1] == 'file-check':
                reply.stdout = b'not office\n'
            return reply
        self.external = reject
        code, result = self.run_reader()
        self.assertEqual((code, result['error']), (1, 'office_validation_failed'))

    def test_permission_denial_does_not_retry_or_leak_details(self):
        self.external = lambda *args, **kwargs: (_ for _ in ()).throw(
            subprocess.CalledProcessError(1, 'host', stderr=b'private credentials'))
        code, result = self.run_reader()
        self.assertEqual((code, result), (1, {'status': 'unread', 'error': 'CalledProcessError'}))

    def test_oversized_output_is_not_silently_truncated(self):
        book = load_workbook(self.path)
        book.active['A2'] = 'x' * 25000
        book.save(self.path)
        book.close()
        self.raw = self.path.read_bytes()
        self.context['inputs'][0]['sha256'] = hashlib.sha256(self.raw).hexdigest()
        code, result = self.run_reader()
        self.assertEqual((code, result['error']), (1, 'readback_too_large_reduce_rows'))
        self.assertNotIn('rows', result)


if __name__ == '__main__':
    unittest.main()
