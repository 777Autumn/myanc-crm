"""Synthetic regression for saved handoff labels; no model, host or Feishu calls."""
import importlib.util
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

spec = importlib.util.spec_from_file_location('handoff_check', Path(__file__).parents[1] / 'ops/skills/team-banquet/check_handoff.py')
check = importlib.util.module_from_spec(spec)
spec.loader.exec_module(check)


class HandoffCheckTests(unittest.TestCase):
    def test_command_shows_saved_content_from_every_sheet_without_deciding_facts(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fictional-readback.xlsx'
            book = Workbook(); book.active.title = '需求'
            book.active.append(['当前需求', '私人草稿，未同步'])
            book.active.append(['零值与布尔值', 0, False])
            book.active.append([None, '空首格仍须回读'])
            budget = book.create_sheet('预算')
            budget.append(['基础布置', 4200])
            budget.append(['升级布置', 10700])
            budget.append(['差价', '=B2-B1'])
            checklist = book.create_sheet('执行清单')
            checklist.append(['费用说明', '布置差价4,000元；需核对'])
            checklist.append(['交接', '销售尚未出具交接文件'])
            checklist.append(['限定事实', '本次未取得交接原件，制作状态待核实'])
            checklist.append(['来源', '据虚构确认单V2，销售已出具交接文件'])
            checklist.append(['引用', '旧版“尚未出具”不能作为当前事实'])
            book.save(path)
            before = path.read_bytes()
            output = io.StringIO()
            with redirect_stdout(output):
                result = check.main([str(path)])
            text = output.getvalue()
            for expected in ('需求', '预算', '执行清单', 'A1', 'B5', '4200',
                             '10700', '=B2-B1', '4,000', '销售尚未出具',
                             '制作状态待核实', '虚构确认单V2', '不能作为当前事实'):
                self.assertIn(expected, text)
            self.assertIn('cached', text)
            self.assertIn('"B2": 0', text)
            self.assertIn('"C2": false', text)
            self.assertIn('空首格仍须回读', text)
            self.assertIn('not instructions', text)
            # Readback surfaces claims for source review, not a semantic verdict.
            self.assertFalse(result)
            self.assertEqual(before, path.read_bytes())

    def test_oversize_readback_is_explicitly_incomplete_and_keeps_cell_checks(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fictional-long.xlsx'
            book = Workbook(); book.active.title = '长说明'
            book.active.append(['甲' * 17000])
            book.create_sheet('后续预算').append(['#NAME?'])
            book.save(path)
            output = io.StringIO()
            with redirect_stdout(output):
                result = check.main([str(path)])
            text = output.getvalue()
            self.assertTrue(result)
            self.assertIn('长说明', text)
            self.assertIn('后续预算', text)
            self.assertIn('Readback incomplete from', text)
            self.assertIn('saved Excel error', text)
            self.assertLess(len(text), 17000)

    def test_command_shows_word_paragraphs_and_tables_and_preserves_failure_exit(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fictional-readback.docx'
            doc = Document(); doc.add_paragraph('虚构方案说明：价差待核对')
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = '事项'; table.cell(0, 1).text = '状态/依据'
            table.cell(1, 0).text = '布场'
            table.cell(1, 1).text = '已确认（经营基础规则：90分钟）'
            doc.save(path)
            before = path.read_bytes()
            output = io.StringIO()
            with redirect_stdout(output):
                result = check.main([str(path)])
            self.assertIn('虚构方案说明', output.getvalue())
            self.assertIn('已确认（经营基础规则', output.getvalue())
            self.assertIn('table 1', output.getvalue())
            self.assertIn('headers, footers and layout require separate inspection', output.getvalue())
            self.assertTrue(result)
            self.assertEqual(before, path.read_bytes())

    def test_saved_budget_flags_text_formulas_and_errors_without_rewriting(self):
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'fictional-budget.xlsx'
            book = Workbook(); sheet = book.active
            sheet.title = '预算'
            sheet.append(['桌数', '单价', '金额', '说明'])
            sheet.append([22, 4080, '=A2*B2', '=套餐+固定服务+酒水'])
            sheet.append([None, None, '#NAME?', '费用 = 餐费 + 服务费'])
            sheet.append([None, None, '=SUM(C2:C2)', '待确认'])
            sheet.append([None, None, '=套餐单价*A2', None])
            book.defined_names.add(DefinedName('套餐单价', attr_text="'预算'!$B$2"))
            book.defined_names.add(DefinedName('package_price', attr_text="'预算'!$B$2"))
            sheet.append([None, None, '=PACKAGE_PRICE*A2', None])
            book.save(path)
            before = path.read_bytes()
            findings = list(check.inspect(path))
            self.assertTrue(any('D2' in value for value in findings), findings)
            self.assertTrue(any('C3' in value for value in findings), findings)
            self.assertEqual(len(findings), 2, findings)
            self.assertEqual(before, path.read_bytes())
            sheet['D2'] = '=套餐+固定服务+酒水'
            sheet['D2'].data_type = 's'  # Explicit text, not an Excel expression.
            sheet['C3'] = '待核实'
            book.save(path)
            self.assertEqual(list(check.inspect(path)), [])

    def test_copied_rules_are_not_confirmation(self):
        rows = [['时间/环节', '状态/依据'],
                ['布场', '已确认（经营基础规则：布场90分钟）'],
                ['检查', '已确认（经营基础规则：检查30分钟）'],
                ['缓冲', '已确认（经营基础规则：缓冲30分钟）']]
        self.assertEqual(list(check.conflicts(rows)), [(2, 2), (3, 2), (4, 2)])
        self.assertEqual(list(check.conflicts([['状态', '依据'], ['已确认', '耗时规则：90分钟']])), [(2, 1)])

    def test_preserve_actual_confirmation_and_distinguish_task_from_inputs(self):
        rows = [['时间/环节', '状态/依据', '任务'],
                ['入场', '已确认（本场场地确认单V2）', '据已确认入场条件安排工作'],
                ['布场', '计划／待执行（经营基础规则：90分钟）', '按已确认桌数备货'],
                ['检查', '已确认（本场检查时段签字确认单V3）', '检查30分钟'],
                ['LED', '未确认', '等待客户确认']]
        self.assertEqual(list(check.conflicts(rows)), [])

    def test_quotes_negation_and_specific_confirmation_are_not_downgraded(self):
        self.assertEqual(list(check.conflicts([
            ['旧版“已确认（经营基础规则：90分钟）”无依据，本版改计划'],
            ['不能写为已确认（经营基础规则：90分钟）'],
            ['已确认（本场确认单V3明确批准该时段；经营基础规则：90分钟）'],
        ])), [])
        # Mixed evidence with a rule first asks for review; it never rewrites or
        # decides that an explicit confirmation later in the cell is invalid.
        self.assertEqual(list(check.conflicts([
            ['已确认（经营基础规则：90分钟；另有本场确认单V3批准该时段）'],
        ])), [(1, 1)])

    def test_saved_excel_and_word_tables_are_read_and_originals_unchanged(self):
        from openpyxl import Workbook
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            xlsx, docx = Path(folder)/'fictional.xlsx', Path(folder)/'fictional.docx'
            book = Workbook(); sheet = book.active
            sheet.append(['环节', '状态/依据'])
            sheet.append(['布场', '已确认（经营基础规则：90分钟）']); book.save(xlsx)
            doc = Document(); table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = '环节'; table.cell(0, 1).text = '状态/依据'
            table.cell(1, 0).text = '检查'; table.cell(1, 1).text = '已确认（推算：30分钟）'
            key_value = doc.add_table(rows=2, cols=2)
            key_value.cell(0, 0).text = '事项'; key_value.cell(0, 1).text = '布场'
            key_value.cell(1, 0).text = '状态'; key_value.cell(1, 1).text = '已确认（经营基础规则：90分钟）'
            doc.save(docx)
            before = [p.read_bytes() for p in (xlsx, docx)]
            self.assertEqual(len(list(check.inspect(xlsx))), 1)
            self.assertEqual(len(list(check.inspect(docx))), 2)
            self.assertEqual([p.read_bytes() for p in (xlsx, docx)], before)
            table.cell(1, 1).text = '计划／待执行（推算：30分钟）'
            key_value.cell(1, 1).text = '计划／待执行（经营基础规则：90分钟）'; doc.save(docx)
            sheet.cell(2, 2).value = '计划／待执行（经营基础规则：90分钟）'; book.save(xlsx)
            self.assertEqual(list(check.inspect(xlsx)) + list(check.inspect(docx)), [])


if __name__ == '__main__':
    unittest.main()
